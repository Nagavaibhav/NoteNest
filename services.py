"""Local extraction and deterministic classification helpers for NoteNest."""

from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader


SUBJECT_RULES = {
    "Computer Science": {
        "programming", "python", "java", "javascript", "algorithm", "data structure",
        "database", "sql", "software engineering", "operating system", "network", "api",
        "compiler", "machine learning", "artificial intelligence", "html", "css",
    },
    "Mathematics": {
        "calculus", "algebra", "derivative", "integral", "matrix", "probability",
        "statistics", "theorem", "equation", "vector", "linear programming",
    },
    "Physics": {
        "physics", "mechanics", "kinematics", "thermodynamics", "electricity", "magnetic",
        "quantum", "wave", "force", "velocity", "energy",
    },
    "Electronics": {
        "circuit", "semiconductor", "transistor", "diode", "microcontroller", "voltage",
        "current", "digital logic", "signal", "embedded",
    },
    "General Studies": {
        "history", "economics", "psychology", "management", "communication", "ethics",
        "environment", "society", "research methodology",
    },
}

TOPIC_RULES = {
    "Data Structures": {"data structure", "linked list", "stack", "queue", "tree", "graph", "hash table"},
    "Algorithms": {"algorithm", "sorting", "searching", "dynamic programming", "complexity", "recursion"},
    "Databases": {"database", "sql", "normalization", "transaction", "erd", "relational"},
    "Web Development": {"html", "css", "javascript", "flask", "api", "frontend", "backend"},
    "Software Engineering": {"software engineering", "requirements", "uml", "testing", "agile", "sdlc"},
    "Calculus": {"calculus", "derivative", "integral", "differential", "limit"},
    "Linear Algebra": {"matrix", "vector", "eigenvalue", "linear algebra"},
    "Probability & Statistics": {"probability", "statistics", "distribution", "hypothesis", "regression"},
    "Mechanics": {"mechanics", "kinematics", "force", "velocity", "acceleration"},
    "Circuits": {"circuit", "voltage", "current", "resistor", "capacitor", "transistor"},
}

STOP_WORDS = {
    "about", "after", "also", "and", "are", "been", "being", "between", "can", "class", "course", "each", "for", "from", "have", "into", "its", "more", "note", "notes", "not", "only", "other", "our", "out", "over", "should", "such", "than", "that", "the", "their", "there", "these", "this", "through", "using", "was", "were", "what", "when", "which", "will", "with", "your",
}


def extract_text(path: Path, extension: str) -> str:
    """Read supported note files locally and return clean text."""
    extension = extension.lower()
    if extension == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    if extension == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if extension == ".docx":
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if extension in {".ppt", ".pptx"}:
        # python-pptx reads Office XML. Convert a legacy PowerPoint to a temporary
        # PPTX with the locally installed Office converter, never an external service.
        presentation_path = path
        if extension == ".ppt":
            with tempfile.TemporaryDirectory() as output_dir:
                completed = subprocess.run(
                    ["soffice", "--headless", "--convert-to", "pptx", "--outdir", output_dir, str(path)],
                    capture_output=True,
                    text=True,
                    timeout=45,
                    check=False,
                )
                converted = Path(output_dir) / f"{path.stem}.pptx"
                if completed.returncode != 0 or not converted.exists():
                    raise ValueError("The PowerPoint file could not be converted for text extraction.")
                presentation = Presentation(str(converted))
                return "\n".join(
                    shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text")
                )
        presentation = Presentation(str(presentation_path))
        return "\n".join(
            shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text")
        )
    raise ValueError("This file type is not supported.")


def _count_matches(text: str, keywords: set[str]) -> int:
    return sum(len(re.findall(r"\b" + re.escape(keyword) + r"\b", text)) for keyword in keywords)


def classify(text: str, title: str = "") -> tuple[str, list[str]]:
    """Return the strongest subject and up to three keyword-derived topics."""
    searchable = f"{title}\n{text}".lower()
    scored_subjects = [(subject, _count_matches(searchable, keywords)) for subject, keywords in SUBJECT_RULES.items()]
    subject, score = max(scored_subjects, key=lambda item: item[1])
    if score == 0:
        subject = "Uncategorized"

    scored_topics = [(topic, _count_matches(searchable, keywords)) for topic, keywords in TOPIC_RULES.items()]
    topics = [topic for topic, score in sorted(scored_topics, key=lambda item: (-item[1], item[0])) if score > 0][:3]
    return subject, topics or ["General notes"]


def extract_topics(text: str, title: str = "") -> list[str]:
    """Create useful topic labels from known study terms and frequent local terms."""
    _, rule_topics = classify(text, title)
    searchable = f"{title}\n{text}".lower()
    words = re.findall(r"[a-z][a-z0-9-]{3,}", searchable)
    frequency: dict[str, int] = {}
    for word in words:
        if word not in STOP_WORDS:
            frequency[word] = frequency.get(word, 0) + 1
    automatic = [word.replace("-", " ").title() for word, count in sorted(frequency.items(), key=lambda item: (-item[1], item[0])) if count >= 2]
    topics: list[str] = []
    for topic in rule_topics + automatic:
        if topic.lower() not in {saved.lower() for saved in topics}:
            topics.append(topic)
        if len(topics) == 5:
            break
    return topics or ["General notes"]


def relevance_score(note: dict | object, query: str) -> int:
    """Rank a note using phrase, title, labels, and text matches; entirely locally."""
    terms = [term for term in re.findall(r"[\w-]+", query.lower()) if len(term) > 1]
    if not terms:
        return 0
    title = str(note["original_name"]).lower()
    subject = str(note["subject"]).lower()
    topics = str(note["topics"]).lower()
    body = str(note["extracted_text"]).lower()
    phrase = " ".join(terms)
    score = 0
    for term in terms:
        score += title.count(term) * 24
        score += subject.count(term) * 12
        score += topics.count(term) * 10
        score += min(body.count(term), 12) * 2
    if phrase in title:
        score += 30
    if phrase in body:
        score += 10
    return score


def make_excerpt(text: str, length: int = 220) -> str:
    clean = re.sub(r"\s+", " ", text).strip()
    return clean[:length].rstrip() + ("..." if len(clean) > length else "")
