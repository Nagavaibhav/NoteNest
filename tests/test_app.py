import io
import os
import tempfile
import unittest
from pathlib import Path

from app import create_app
from docx import Document
from pptx import Presentation
from services import classify, extract_text, extract_topics, relevance_score


class NoteNestTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app({
            "TESTING": True,
            "SECRET_KEY": "test-secret",
            "DATABASE": os.path.join(self.temp.name, "test.db"),
            "UPLOAD_FOLDER": os.path.join(self.temp.name, "uploads"),
        })
        self.client = self.app.test_client()

    def tearDown(self):
        self.temp.cleanup()

    def csrf(self):
        with self.client.session_transaction() as session:
            return session["csrf_token"]

    def register_and_login(self):
        self.client.get("/register")
        self.client.post("/register", data={"csrf_token": self.csrf(), "name": "Asha Student", "email": "asha@example.edu", "password": "long-password"})
        self.client.get("/login")
        return self.client.post("/login", data={"csrf_token": self.csrf(), "email": "asha@example.edu", "password": "long-password"}, follow_redirects=True)

    def test_registration_login_and_keyword_organization(self):
        response = self.register_and_login()
        self.assertIn(b"My notes", response.data)
        self.client.get("/upload")
        response = self.client.post("/upload", data={
            "csrf_token": self.csrf(),
            "file": (io.BytesIO(b"Software engineering requirements and testing for a web API."), "review-notes.txt"),
        }, content_type="multipart/form-data", follow_redirects=True)
        self.assertIn(b"Computer Science", response.data)
        self.assertIn(b"Software Engineering", response.data)
        response = self.client.get("/dashboard?q=requirements")
        self.assertIn(b"review-notes.txt", response.data)

    def test_upload_requires_login(self):
        response = self.client.get("/upload", follow_redirects=True)
        self.assertIn(b"<h2>Sign in</h2>", response.data)

    def test_rejects_unsupported_file(self):
        self.register_and_login()
        self.client.get("/upload")
        response = self.client.post("/upload", data={"csrf_token": self.csrf(), "file": (io.BytesIO(b"x"), "notes.exe")}, content_type="multipart/form-data", follow_redirects=True)
        self.assertIn(b"Use a PDF, PPT/PPTX, DOCX, or TXT file", response.data)

    def test_docx_and_pptx_text_extraction(self):
        docx_path = Path(self.temp.name) / "lecture.docx"
        document = Document()
        document.add_paragraph("Calculus derivatives and integrals")
        document.save(docx_path)
        self.assertIn("derivatives", extract_text(docx_path, ".docx"))

        pptx_path = Path(self.temp.name) / "slides.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Database normalization"
        presentation.save(pptx_path)
        self.assertIn("normalization", extract_text(pptx_path, ".pptx"))

    def test_keyword_classification_is_transparent(self):
        subject, topics = classify("SQL database normalization and transactions")
        self.assertEqual(subject, "Computer Science")
        self.assertIn("Databases", topics)

    def test_keyword_matching_respects_word_boundaries(self):
        # "pythonic" must not be miscounted as a match for the keyword "python".
        from services import _count_matches
        self.assertEqual(_count_matches("this code is pythonic", {"python"}), 0)
        self.assertEqual(_count_matches("i love python programming", {"python"}), 1)

    def test_module_bookmark_topic_extraction_and_ranked_search(self):
        self.register_and_login()
        self.client.get("/upload")
        self.client.post("/upload", data={
            "csrf_token": self.csrf(), "module": "Module 3",
            "file": (io.BytesIO(b"Database normalization is covered in this database lecture. SQL joins are included."), "database-notes.txt"),
        }, content_type="multipart/form-data")
        response = self.client.get("/dashboard?module=Module+3&type=.txt&q=database&sort=relevance")
        self.assertIn(b"database-notes.txt", response.data)
        self.assertIn(b"Module 3", response.data)
        response = self.client.post("/notes/1/bookmark", data={"csrf_token": self.csrf()}, follow_redirects=True)
        self.assertIn(b"Bookmark updated", response.data)
        response = self.client.get("/dashboard?bookmarked=1")
        self.assertIn(b"database-notes.txt", response.data)
        self.assertIn("Databases", extract_topics("Database query and database normalization"))
        note = {"original_name": "database-notes.txt", "subject": "Computer Science", "topics": "Databases", "extracted_text": "database database"}
        self.assertGreater(relevance_score(note, "database"), 0)

    def test_export_and_import_preserve_note_metadata(self):
        self.register_and_login()
        self.client.get("/upload")
        self.client.post("/upload", data={
            "csrf_token": self.csrf(), "module": "Unit 1",
            "file": (io.BytesIO(b"Python programming practice notes"), "python.txt"),
        }, content_type="multipart/form-data")
        self.client.post("/notes/1/bookmark", data={"csrf_token": self.csrf()})
        exported = self.client.get("/export")
        self.assertEqual(exported.status_code, 200)
        self.assertTrue(exported.data.startswith(b"PK"))

        second = self.app.test_client()
        second.get("/register")
        with second.session_transaction() as session:
            token = session["csrf_token"]
        second.post("/register", data={"csrf_token": token, "name": "Bala Student", "email": "bala@example.edu", "password": "long-password"})
        second.get("/login")
        with second.session_transaction() as session:
            token = session["csrf_token"]
        second.post("/login", data={"csrf_token": token, "email": "bala@example.edu", "password": "long-password"})
        second.get("/import")
        with second.session_transaction() as session:
            token = session["csrf_token"]
        response = second.post("/import", data={"csrf_token": token, "archive": (io.BytesIO(exported.data), "backup.zip")}, content_type="multipart/form-data", follow_redirects=True)
        self.assertIn(b"Imported 1 note", response.data)
        self.assertIn(b"python.txt", response.data)
        self.assertIn(b"Unit 1", response.data)


if __name__ == "__main__":
    unittest.main()
